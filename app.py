import streamlit as st
import pypdf
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from io import BytesIO
import re
from datetime import datetime, timedelta

st.set_page_config(page_title="Invoice to Excel Converter", page_icon="📑", layout="wide")

st.title("📑 Landscaping Invoice to Excel Converter")
st.write("Upload an invoice PDF to extract line items and export directly into your accounting import spreadsheet.")

uploaded_file = st.file_uploader("Choose an Invoice PDF", type=["pdf"])

def clean_description_block(desc_text):
    """Cleans up the description text block and removes table headers/footers."""
    lines = desc_text.split('\n')
    cleaned = []
    for l in lines:
        s = l.strip()
        if not s:
            continue
        # Drop table header lines
        if "Description Qty / UOM" in s or s == "Description":
            continue
        # Drop footer phone/website lines
        if "713-657-0875" in s or "lasallelandscaping.com" in s:
            continue
        # Strip trailing price from the line if present (e.g. "Plant Installation ... $747.50")
        m_amt = re.search(r'\s+\$([\d,]+\.\d{2})$', s)
        if m_amt:
            s = s[:m_amt.start()].strip()
            if not s:
                continue
        cleaned.append(s)
    return "\n".join(cleaned)

def parse_invoices(pdf_bytes):
    reader = pypdf.PdfReader(BytesIO(pdf_bytes))
    
    # Group pages by invoice (handles multi-page invoices like 1235, 1236, etc.)
    invoices = []
    current_inv = None

    for page in reader.pages:
        text = page.extract_text() or ""
        m = re.search(r'Invoice\s+(\d+)', text)
        if m:
            if current_inv:
                invoices.append(current_inv)
            current_inv = {
                'inv_num': m.group(1),
                'pages': [text]
            }
        else:
            if current_inv:
                current_inv['pages'].append(text)

    if current_inv:
        invoices.append(current_inv)

    records = []

    for inv in invoices:
        full_text = "\n".join(inv['pages'])
        inv_num = inv['inv_num']

        # 1. Invoice Date & Terms -> Due Date
        date_match = re.search(r'Date\s+PO#\s*\n\s*([0-9/]+)', full_text)
        inv_date_str = ""
        due_date_str = ""
        if date_match:
            raw_date = date_match.group(1).strip()
            try:
                dt = datetime.strptime(raw_date, "%m/%d/%y")
                inv_date_str = f"{dt.month}/{dt.day}/{dt.year}"
                if "Due on Receipt" in full_text:
                    due_date_str = f"{dt.month}/{dt.day:02d}/{dt.year}"
                else:  # Net 30
                    due_dt = dt + timedelta(days=30)
                    due_date_str = f"{due_dt.month}/{due_dt.day:02d}/{due_dt.year}"
            except:
                inv_date_str = raw_date
                due_date_str = raw_date

        # 2. Customer: Extract the first line of Property Address
        cust = ""
        bt_idx = full_text.find("Bill To Property Address")
        desc_idx = full_text.find("Description", bt_idx) if bt_idx != -1 else -1
        if bt_idx != -1 and desc_idx != -1:
            addr_block = full_text[bt_idx + len("Bill To Property Address"):desc_idx]
            addr_lines = [l.strip() for l in addr_block.split('\n') if l.strip()]
            # Bill To block finishes at the first City/State/Zip line (e.g. TX 77075)
            # The line right after is the first line of Property Address (the customer name)
            found_zip_idx = -1
            for i, l in enumerate(addr_lines):
                if re.search(r'[A-Z]{2}\s+\d{5}', l):
                    found_zip_idx = i
                    break
            if found_zip_idx != -1 and found_zip_idx + 1 < len(addr_lines):
                cust = addr_lines[found_zip_idx + 1]
            elif addr_lines:
                cust = addr_lines[0]

        # 3. Subtotal / Unit Price
        sub_match = re.search(r'Subtotal\s*\$?([\d,]+\.\d{2})', full_text)
        unit_price = float(sub_match.group(1).replace(',', '')) if sub_match else 0.0

        # 4. Tax
        tax_match = re.search(r'Sales Tax\s*\$?([\d,]+\.\d{2})', full_text)
        has_tax = "No"
        if tax_match:
            t_val = float(tax_match.group(1).replace(',', ''))
            if t_val > 0:
                has_tax = "Yes"

        # 5. Full Description
        desc_start = full_text.find("Description Qty / UOM")
        if desc_start == -1:
            desc_start = full_text.find("Description")
        sub_start = full_text.find("Subtotal", desc_start)
        raw_desc = full_text[desc_start:sub_start] if sub_start != -1 else full_text[desc_start:]
        clean_desc = clean_description_block(raw_desc)

        records.append({
            "Post": "Yes",
            "Invoice Date": inv_date_str,
            "Due Date": due_date_str,
            "Invoice Number": inv_num,
            "Transaction Type": "Invoice",
            "Customer": cust,
            "Vendor": "",
            "Currency Code": "",
            "Products/Services": "Side Jobs",
            "Description": clean_desc,
            "Qty": 1,
            "Discount %": "",
            "Unit Price": unit_price,
            "Category": "",
            "Location": "",
            "Class": "",
            "Tax": has_tax
        })

    return records

def create_excel(records):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Invoice Import"
    
    headers = [
        "Post", "Invoice Date", "Due Date", "Invoice Number", "Transaction Type",
        "Customer", "Vendor", "Currency Code", "Products/Services", "Description",
        "Qty", "Discount %", "Unit Price", "Category", "Location", "Class", "Tax"
    ]
    ws.append(headers)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_idx, r in enumerate(records, start=2):
        row_values = [r[h] for h in headers]
        ws.append(row_values)
        fill_color = "F9FAFC" if row_idx % 2 == 0 else "FFFFFF"
        row_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        
        for col_idx, cell in enumerate(ws[row_idx], start=1):
            cell.fill = row_fill
            cell.border = thin_border
            cell.font = Font(name="Calibri", size=10)
            col_name = headers[col_idx - 1]
            if col_name in ["Post", "Transaction Type", "Qty", "Tax", "Invoice Date", "Due Date", "Invoice Number"]:
                cell.alignment = Alignment(horizontal="center", vertical="top")
            elif col_name == "Unit Price":
                cell.number_format = '$#,##0.00'
                cell.alignment = Alignment(horizontal="right", vertical="top")
            elif col_name == "Description":
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            else:
                cell.alignment = Alignment(horizontal="left", vertical="top")

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or '').split('\n')[0]) for cell in col)
        ws.column_dimensions[col_letter].width = max(min(max_len + 4, 42), 12)
    ws.column_dimensions['J'].width = 50

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output

if uploaded_file is not None:
    with st.spinner("Processing PDF and extracting line items..."):
        data = parse_invoices(uploaded_file.read())
        
    if data:
        st.success(f"Successfully processed {len(data)} invoices!")
        excel_data = create_excel(data)
        
        st.download_button(
            label="📥 Download Excel Spreadsheet",
            data=excel_data,
            file_name="Extracted_Invoices.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
    else:
        st.error("No invoice data found in the uploaded file.")
